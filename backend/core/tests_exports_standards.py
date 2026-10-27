"""Feature 015 exported-file regressions, using only installed app dependencies."""
import base64
import io
import re
import zlib
from types import SimpleNamespace
from datetime import date, datetime
from unittest.mock import patch

from django.test import TestCase
from openpyxl import load_workbook

from accounts.models import Department
from assets.models import AssetTransfer
from assets.tests_b6_conformance import _asset, _catalog
from .exports import build_excel, build_table_pdf, build_transfer_pdf
from .models import Setting
from .tests_b8_conformance import _client


def pdf_text(data):
    """Read ReportLab's actual PDF text streams without another PDF dependency.

    These fixtures use the standard ASCII fonts; PDF string escapes, ASCII85 and
    Flate compression are decoded rather than disabling normal export compression.
    """
    chunks = []
    for stream in re.findall(rb'stream\r?\n(.*?)endstream', data, re.S):
        stream = stream.strip()
        if stream.endswith(b'~>'):
            stream = base64.a85decode(b'<~' + stream, adobe=True)
        try:
            stream = zlib.decompress(stream)
        except zlib.error:
            pass
        for token in re.findall(rb'\(((?:\\.|[^\\)])*)\)\s*Tj', stream):
            token = re.sub(rb'\\([0-7]{1,3})', lambda m: bytes([int(m[1], 8)]), token)
            token = re.sub(rb'\\([()\\])', rb'\1', token)
            chunks.append(token.decode('cp1252'))
    return ' '.join(chunks)


def delivered(response):
    return b''.join(response.streaming_content)


class LiteralPdfTests(TestCase):
    def assert_literal(self, data, values):
        self.assertTrue(data.startswith(b'%PDF'))
        text = ''.join(pdf_text(data).split())
        for value in values:
            self.assertIn(''.join(value.split()), text)
        self.assertIn('AuthorizedBy' if 'AssetCustodyTransferDocument' in text else 'PreparedBy', text)

    def test_shared_pdf_preserves_company_title_headers_and_cells(self):
        values = ['<b>broken', 'A & B', '<i>literal</i>', '<img src="https://example.invalid/probe.png"/>']
        for position in range(4):
            with self.subTest(position=position), patch('reportlab.platypus.paraparser.ImageReader', side_effect=AssertionError('resource fetch')) as reader:
                company, title, header, cell = values[position:] + values[:position]
                result = build_table_pdf(title, [header], [[cell]], company_name=company)
                self.assert_literal(result.getvalue(), values)
                reader.assert_not_called()

    def test_transfer_detail_values_are_literal_in_every_user_text_position(self):
        values = {key: f'<b>{key} & value' for key in (
            'status', 'asset_id', 'name', 'tag_number', 'serial_number', 'category',
            'condition', 'previous_owner', 'previous_department', 'new_owner',
            'new_department', 'requested_by', 'approved_by', 'remarks')}
        asset = SimpleNamespace(**{key: values[key] for key in ('asset_id', 'name', 'tag_number', 'serial_number')},
                                category=SimpleNamespace(category_name=values['category']))
        transfer = SimpleNamespace(id=123, asset=asset, transfer_date=None,
                                   **{key: values[key] for key in ('status', 'condition', 'remarks')},
                                   **{key: SimpleNamespace(user_name=values[key]) for key in (
                                       'previous_owner', 'new_owner', 'requested_by', 'approved_by')},
                                   **{key: SimpleNamespace(name=values[key]) for key in (
                                       'previous_department', 'new_department')})
        with patch('reportlab.platypus.paraparser.ImageReader', side_effect=AssertionError('resource fetch')) as reader:
            self.assert_literal(build_transfer_pdf(transfer).getvalue(), list(values.values()))
            reader.assert_not_called()

    def test_public_report_and_transfer_deliver_literal_text_and_preserve_permissions(self):
        client, profile = _client('literal-exporter', ('report_access', 'asset_view'))
        cat = _catalog()
        company = '<b>Company & Co'
        name = '<b>broken'
        category = '<i>Category</i>'
        remark = '<img src="https://example.invalid/probe.png"/>'
        cat['category'].category_name = category
        cat['category'].save()
        asset = _asset('PDF-LITERAL', cat, name=name)
        department = Department.objects.create(name='A & B')
        transfer = AssetTransfer.objects.create(asset=asset, new_owner=profile,
                                                new_department=department, remarks=remark)
        Setting.set('company_name', company)
        for url, values in [('/reports/register/export/pdf/', [company, name, category]),
                            (f'/transfer/{transfer.pk}/document/', [company, name, category, remark, 'A & B'])]:
            with self.subTest(url=url), patch('reportlab.platypus.paraparser.ImageReader', side_effect=AssertionError('resource fetch')) as reader:
                response = client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assert_literal(delivered(response), values)
                reader.assert_not_called()
        denied, _ = _client('literal-denied', ())
        self.assertEqual(denied.get('/reports/register/export/pdf/').status_code, 403)
        self.assertEqual(denied.get(f'/transfer/{transfer.pk}/document/').status_code, 403)


class LiteralExcelTests(TestCase):
    def test_shared_workbook_preserves_literal_strings_and_nonstring_types(self):
        strings = ['=1+1', '+x', '-x', '@x']
        typed = [0, -1, True, False, date(2026, 10, 27), datetime(2026, 10, 27, 12, 30), None]
        rows = [strings + typed]
        workbook = load_workbook(build_excel('=Title', strings + list(range(len(typed))), rows, company_name='=Company'))
        sheet = workbook.active
        self.assertEqual(sheet.cell(1, 1).value, '=Company — =Title')
        self.assertEqual(sheet.cell(1, 1).data_type, 's')
        for index, text in enumerate(strings, start=1):
            for row in (2, 3):
                self.assertEqual(sheet.cell(row, index).value, text)
                self.assertEqual(sheet.cell(row, index).data_type, 's')
        for index, value in enumerate(typed, start=5):
            cell = sheet.cell(3, index)
            expected = datetime.combine(value, datetime.min.time()) if type(value) is date else value
            self.assertEqual(cell.value, expected)
            self.assertEqual(cell.data_type, 'b' if isinstance(value, bool) else 'd' if isinstance(value, (date, datetime)) else 'n')

    def test_public_delivered_workbook_preserves_strings_order_and_permission(self):
        client, _ = _client('excel-exporter', ('report_access',))
        cat = _catalog()
        strings = ['=1+1', '+x', '-x', '@x']
        for index, value in enumerate(strings):
            _asset(f'XLSX-{index}', cat, name=value)
        Setting.set('company_name', '=Company')
        response = client.get('/reports/register/export/excel/')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(delivered(response))).active
        self.assertEqual(sheet.cell(1, 1).value, '=Company — Asset Register Report')
        self.assertEqual(sheet.cell(1, 1).data_type, 's')
        headers = [cell.value for cell in sheet[2]]
        name_index = headers.index('Name') + 1
        names = [sheet.cell(row, name_index) for row in range(3, 7)]
        self.assertEqual([cell.value for cell in names], strings)
        self.assertTrue(all(cell.data_type == 's' for cell in names))
        denied, _ = _client('excel-denied', ())
        self.assertEqual(denied.get('/reports/register/export/excel/').status_code, 403)
