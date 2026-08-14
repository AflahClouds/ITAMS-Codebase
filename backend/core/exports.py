"""Server-side PDF (reportlab) and Excel (openpyxl) builders."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

BRAND_BLUE = colors.HexColor("#01014b")


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.grey)
    canvas.drawString(15 * mm, 8 * mm,
                      f"Generated on {datetime.now().strftime('%d-%b-%Y %H:%M')}")
    canvas.drawRightString(doc.pagesize[0] - 15 * mm, 8 * mm,
                           f"Page {canvas.getPageNumber()}")
    canvas.restoreState()


def build_table_pdf(title, headers, rows, company_name="Cloud Shoe Factory"):
    """Landscape A4 report PDF with a banner header and striped table."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Banner", parent=styles["Title"],
                                 textColor=colors.white, fontSize=14, leading=18)
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=7, leading=9)

    banner = Table([[Paragraph(f"{company_name} — {title}", title_style)]],
                   colWidths=[doc.width])
    banner.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), BRAND_BLUE),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
    ]))

    data = [[Paragraph(f"<b>{h}</b>", cell_style) for h in headers]]
    for row in rows:
        data.append([Paragraph(str(cell if cell is not None else "-"), cell_style)
                     for cell in row])

    table = Table(data, repeatRows=1, colWidths=[doc.width / len(headers)] * len(headers))
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eaf6")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c5cae9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f6fa")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))

    sig = Table([["", ""], ["Prepared By", "Approved By"]],
                colWidths=[doc.width / 2] * 2, rowHeights=[14 * mm, 6 * mm])
    sig.setStyle(TableStyle([
        ("LINEABOVE", (0, 1), (0, 1), 0.5, colors.black),
        ("LINEABOVE", (1, 1), (1, 1), 0.5, colors.black),
        ("FONTSIZE", (0, 1), (-1, 1), 8),
        ("ALIGN", (0, 1), (-1, 1), "CENTER"),
    ]))

    doc.build([banner, Spacer(1, 6 * mm), table, Spacer(1, 12 * mm), sig],
              onFirstPage=_footer, onLaterPages=_footer)
    buf.seek(0)
    return buf


def build_transfer_pdf(transfer, company_name="Cloud Shoe Factory"):
    """Portrait A4 custody transfer document with a signature section."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=15 * mm, bottomMargin=18 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Banner", parent=styles["Title"],
                                 textColor=colors.white, fontSize=14, leading=18)
    label = ParagraphStyle("Label", parent=styles["Normal"], fontSize=9,
                           textColor=colors.HexColor("#555555"))
    value = ParagraphStyle("Value", parent=styles["Normal"], fontSize=10)

    asset = transfer.asset
    banner = Table([[Paragraph(f"{company_name}<br/>Asset Custody Transfer Document",
                               title_style)]], colWidths=[doc.width])
    banner.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), BRAND_BLUE),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 12),
    ]))

    def row(name, val):
        return [Paragraph(name, label), Paragraph(str(val or "-"), value)]

    details = Table([
        row("Transfer ID", f"TRF-{transfer.id:05d}"),
        row("Status", transfer.status),
        row("Date", transfer.transfer_date.strftime("%d-%b-%Y %H:%M") if transfer.transfer_date else "-"),
        row("Asset ID", asset.asset_id),
        row("Asset Name", asset.name or "-"),
        row("Tag Number", asset.tag_number),
        row("Serial Number", asset.serial_number),
        row("Category", asset.category.category_name if asset.category else "-"),
        row("Condition at Transfer", transfer.condition or "-"),
        row("Previous Owner", transfer.previous_owner.user_name if transfer.previous_owner else "IT Store"),
        row("Previous Department", transfer.previous_department.name if transfer.previous_department else "-"),
        row("New Owner", transfer.new_owner.user_name if transfer.new_owner else "-"),
        row("New Department", transfer.new_department.name if transfer.new_department else "-"),
        row("Requested By", transfer.requested_by.user_name if transfer.requested_by else "-"),
        row("Approved By", transfer.approved_by.user_name if transfer.approved_by else "-"),
        row("Remarks", transfer.remarks or "-"),
    ], colWidths=[doc.width * 0.35, doc.width * 0.65])
    details.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c5cae9")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f5f6fa")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))

    sig = Table([["", "", ""],
                 ["Released By\n(Previous Owner)", "Received By\n(New Owner)", "Authorized By\n(IT Department)"]],
                colWidths=[doc.width / 3] * 3, rowHeights=[20 * mm, 10 * mm])
    sig.setStyle(TableStyle([
        ("LINEABOVE", (0, 1), (0, 1), 0.5, colors.black),
        ("LINEABOVE", (1, 1), (1, 1), 0.5, colors.black),
        ("LINEABOVE", (2, 1), (2, 1), 0.5, colors.black),
        ("FONTSIZE", (0, 1), (-1, 1), 8),
        ("ALIGN", (0, 1), (-1, 1), "CENTER"),
    ]))

    doc.build([banner, Spacer(1, 8 * mm), details, Spacer(1, 20 * mm), sig],
              onFirstPage=_footer, onLaterPages=_footer)
    buf.seek(0)
    return buf


def build_excel(title, headers, rows, company_name="Cloud Shoe Factory"):
    """Branded xlsx workbook with banner, styled headers, auto column widths."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
    banner = ws.cell(row=1, column=1, value=f"{company_name} — {title}")
    banner.font = Font(bold=True, size=13, color="FFFFFF")
    banner.fill = PatternFill("solid", fgColor="01014B")
    banner.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    header_fill = PatternFill("solid", fgColor="E8EAF6")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for r, row in enumerate(rows, start=3):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)

    for col in range(1, len(headers) + 1):
        longest = len(str(headers[col - 1]))
        for row in rows:
            if col <= len(row):
                longest = max(longest, len(str(row[col - 1] or "")))
        ws.column_dimensions[get_column_letter(col)].width = min(longest + 3, 45)

    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_csv(title, headers, rows, company_name="Cloud Shoe Factory"):
    """CSV of the same rows the PDF and Excel exports carry.

    B9 ALR-001 and ALR-005 require every export to enforce its access decision
    server-side and to contain exactly what the requester is authorized to read.
    Generating CSV in the browser from data already on screen satisfied neither: it
    never passed through an authorization check and produced no audit entry, so the one
    format most likely to be mailed onward was the one with no record of who took it.
    """
    import csv
    import io as _io

    text = _io.StringIO()
    writer = csv.writer(text, lineterminator="\n")
    writer.writerow([company_name])
    writer.writerow([title])
    writer.writerow([])
    writer.writerow(list(headers))
    for row in rows:
        writer.writerow(["" if cell is None else cell for cell in row])
    # BOM so Excel opens the file as UTF-8 rather than guessing a codepage.
    return _io.BytesIO(text.getvalue().encode("utf-8-sig"))
